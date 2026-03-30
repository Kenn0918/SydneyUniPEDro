"""
RAGflow Batch Labeling Script
------------------------------
For each row in an Excel file, sends the Title + Abstract to a RAGflow chat,
extracts the label (0 or 1) from the response, and writes them back.

Columns:
  B - Title
  C - Abstract
  D - Final label  (0 or 1, written by this script)
  E - Timestamp    (written by this script)
"""

import os
import re
import time
import requests
import openpyxl
from datetime import datetime

# ── Configuration ──────────────────────────────────────────────────────────────
RAGFLOW_BASE_URL = ""   # e.g. "http://192.168.1.100" or "https://your-ragflow-host"
API_KEY          = ""
CHAT_ID          = ""
EXCEL_PATH       = ""

# How long to wait between requests (seconds) to avoid rate-limiting
REQUEST_DELAY = 2

REQUEST_TIMEOUT = 15   # seconds per request
MAX_RETRIES     = 3    # retry attempts before skipping a row

# ── Helpers ────────────────────────────────────────────────────────────────────

def create_session(row_idx: int) -> str:
    """Create a fresh session for this record and return its session_id."""
    url = f"{RAGFLOW_BASE_URL}/api/v1/chats/{CHAT_ID}/sessions"
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json",
    }
    resp = requests.post(url, json={"name": f"row_{row_idx}"}, headers=headers, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    if data.get("code") != 0:
        raise RuntimeError(f"Failed to create session: {data}")
    return data["data"]["id"]


def ask_ragflow(title: str, abstract: str, session_id: str) -> int:
    """Send one record to RAGflow within a fresh session and return the label (0 or 1)."""
    url = f"{RAGFLOW_BASE_URL}/api/v1/chats/{CHAT_ID}/completions"
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json",
    }
    question = f"Title: {title}\n\nAbstract: {abstract}"
    payload = {
        "question": question,
        "session_id": session_id,
        "stream": False,
    }

    response = requests.post(url, json=payload, headers=headers, timeout=REQUEST_TIMEOUT)
    response.raise_for_status()
    data = response.json()

    if data.get("code") != 0:
        raise RuntimeError(f"RAGflow error: {data}")

    answer = data["data"]["answer"].strip()

    # Response is just "0" or "1"
    m = re.search(r"\b([01])\b", answer)
    if m:
        return int(m.group(1))
    raise RuntimeError(f"Could not parse label from response: {answer!r}")


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    total = ws.max_row - 1  # exclude header
    processed = 0
    skipped = 0

    print(f"Found {total} records (excluding header). Starting...\n")

    for row_idx in range(2, ws.max_row + 1):
        title    = ws.cell(row=row_idx, column=2).value  # B
        abstract = ws.cell(row=row_idx, column=3).value  # C
        existing = ws.cell(row=row_idx, column=4).value  # D (Final label)

        # Skip rows that are already labeled
        if existing is not None and str(existing).strip() != "":
            skipped += 1
            print(f"Row {row_idx}: already labeled ({existing}), skipping.")
            continue

        if not title and not abstract:
            print(f"Row {row_idx}: empty row, skipping.")
            continue

        title    = str(title or "").strip()
        abstract = str(abstract or "").strip()

        print(f"Row {row_idx}: processing... ", end="", flush=True)

        last_error = None
        label = None
        for attempt in range(1, MAX_RETRIES + 1):
            try:
                session_id = create_session(row_idx)
                label = ask_ragflow(title, abstract, session_id)
                break  # success — exit retry loop
            except requests.exceptions.Timeout:
                last_error = f"Timeout (attempt {attempt}/{MAX_RETRIES})"
                print(f"\n  {last_error}, retrying..." if attempt < MAX_RETRIES else f"\n  {last_error}, giving up.", end="", flush=True)
            except Exception as e:
                last_error = str(e)
                print(f"\n  Error (attempt {attempt}/{MAX_RETRIES}): {last_error}, retrying..." if attempt < MAX_RETRIES else f"\n  Error (attempt {attempt}/{MAX_RETRIES}): {last_error}, giving up.", end="", flush=True)

        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if label is not None:
            ws.cell(row=row_idx, column=4).value = label  # D
            ws.cell(row=row_idx, column=5).value = ts     # E
            processed += 1
            print(f"  label={label}  [{ts}]")
        else:
            ws.cell(row=row_idx, column=4).value = "ERROR"
            ws.cell(row=row_idx, column=5).value = ts
            print(f"  SKIPPED [{ts}]")

        # Save to temp file first, then replace — prevents corruption on Ctrl+C
        tmp = EXCEL_PATH + ".tmp"
        wb.save(tmp)
        os.replace(tmp, EXCEL_PATH)

        time.sleep(REQUEST_DELAY)

    print(f"\nDone. Processed: {processed}, Skipped (already labeled): {skipped}.")


if __name__ == "__main__":
    main()
